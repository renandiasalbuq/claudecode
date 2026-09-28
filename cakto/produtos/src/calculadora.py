"""Gera a planilha 'Calculadora do Seu Número' (.xlsx).

Uso: python3 calculadora.py  ->  ../entregaveis/02-Calculadora-do-Seu-Numero.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).resolve().parent.parent / "entregaveis" / "02-Calculadora-do-Seu-Numero.xlsx"

VERDE, VERDE3, OURO, OURO_ESC = "0E2B25", "1E4A3D", "C9A45C", "8A6A2E"
CREME, PAPEL, ENTRADA, LINHA = "F3EDE2", "FBF8F2", "FFF6DB", "D8CFBE"
BRL = '"R$" #,##0.00'

fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
fina = Side(style="thin", color=LINHA)
borda = Border(left=fina, right=fina, top=fina, bottom=fina)
ouro_borda = Border(*(Side(style="medium", color=OURO),) * 4)

wb = Workbook()


def base(ws, titulo, subtitulo, larguras):
    ws.sheet_view.showGridLines = False
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 12
    ws["B2"] = titulo
    ws["B2"].font = Font(name="Georgia", size=20, bold=True, color=VERDE)
    ws["B3"] = subtitulo
    ws["B3"].font = Font(size=10.5, color="5B645F")
    ws.row_dimensions[2].height = 32


def rotulo(c, texto, cor=VERDE, negrito=True, tam=10.5):
    c.value = texto
    c.font = Font(size=tam, bold=negrito, color=cor)
    c.alignment = Alignment(vertical="center", wrap_text=True)


def entrada(c, valor=None, fmt=BRL):
    c.value = valor
    c.fill = fill(ENTRADA)
    c.border = ouro_borda
    c.number_format = fmt
    c.font = Font(size=11, bold=True, color=VERDE)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.protection = Protection(locked=False)


def saida(c, formula, fmt=BRL, destaque=False):
    c.value = formula
    c.number_format = fmt
    c.border = borda
    c.alignment = Alignment(horizontal="right", vertical="center")
    if destaque:
        c.fill = fill(VERDE)
        c.font = Font(size=13, bold=True, color="FFFFFF")
    else:
        c.fill = fill(PAPEL)
        c.font = Font(size=11, bold=True, color=VERDE)


def faixa(ws, linha, texto, cols="B:E"):
    a, b = cols.split(":")
    ws.merge_cells(f"{a}{linha}:{b}{linha}")
    c = ws[f"{a}{linha}"]
    c.value = texto
    c.fill = fill(VERDE)
    c.font = Font(size=9.5, bold=True, color="E3C98F")
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[linha].height = 20


# ---------------------------------------------------------------- Comece aqui
ws = wb.active
ws.title = "Comece aqui"
base(ws, "Calculadora do Seu Número", "Descubra quanto é a sua reserva e quanto guardar por mês para chegar lá.",
     {"A": 3, "B": 6, "C": 84})
passos = [
    ("1", "Vá na aba  1 · Despesas  e preencha as células AMARELAS com o valor mensal de cada despesa essencial."),
    ("2", "Vá na aba  2 · Seu Número. Confira quantos meses de reserva você quer (padrão: 6), em quantos meses quer "
          "chegar lá (padrão: 18) e quanto já tem guardado."),
    ("3", "Pronto: a planilha mostra o seu número e quanto guardar por mês e por semana. Mude o prazo e o valor "
          "mensal muda na hora."),
    ("4", "Todo mês, vá na aba  3 · Acompanhamento  e anote quanto depositou. A planilha mostra quanto já juntou, "
          "quanto falta e se você está no ritmo."),
    ("5", "A aba  4 · Simulador de prazo  compara quanto guardar por mês em prazos de 6 a 36 meses."),
]
r = 5
for n, t in passos:
    c = ws.cell(r, 2, n)
    c.font = Font(name="Georgia", size=16, bold=True, color=OURO)
    c.alignment = Alignment(horizontal="center", vertical="top")
    rotulo(ws.cell(r, 3), t, negrito=False, cor="1B2622", tam=11)
    ws.row_dimensions[r].height = 36
    r += 1
r += 1
rotulo(ws.cell(r, 3), "Só preencha as células AMARELAS. As outras têm fórmulas e estão protegidas para você não apagar sem querer.",
       cor=OURO_ESC)
ws.cell(r, 3).fill = fill(ENTRADA)
ws.row_dimensions[r].height = 30
r += 2
rotulo(ws.cell(r, 3), "Funciona no Excel, no Google Planilhas (Arquivo → Importar) e no LibreOffice. "
       "No celular, use o app do Google Planilhas ou do Excel.", negrito=False, cor="5B645F", tam=9.5)
ws.row_dimensions[r].height = 28
r += 1
rotulo(ws.cell(r, 3), "Conteúdo educativo. Valores sem considerar rendimento; confira as regras vigentes do lugar "
       "onde você guarda o dinheiro.", negrito=False, cor="5B645F", tam=9.5)
ws.row_dimensions[r].height = 28

# ---------------------------------------------------------------- 1 · Despesas
wd = wb.create_sheet("1 · Despesas")
base(wd, "Suas despesas essenciais", "Valores MENSAIS. Conta anual? Divida por 12. Use extratos reais dos últimos 3 meses.",
     {"A": 3, "B": 46, "C": 18, "D": 3, "E": 44})
faixa(wd, 5, "DESPESA ESSENCIAL", "B:C")
itens = [
    "Moradia (aluguel ou financiamento, condomínio)",
    "Contas de casa (luz, água, gás)",
    "Internet e celular (plano básico)",
    "Alimentação (mercado, sem delivery)",
    "Transporte (para trabalhar ou procurar trabalho)",
    "Saúde (plano, remédios de uso contínuo)",
    "Educação obrigatória (escola dos filhos)",
    "Dívidas que não podem atrasar (parcelas)",
    "Outro essencial 1",
    "Outro essencial 2",
    "Outro essencial 3",
]
PRIMEIRA = 6
for i, nome in enumerate(itens):
    l = PRIMEIRA + i
    c = wd.cell(l, 2, nome)
    c.font = Font(size=10.5, color="1B2622")
    c.border = borda
    c.alignment = Alignment(vertical="center", indent=1)
    if nome.startswith("Outro"):
        c.fill = fill(ENTRADA)
        c.protection = Protection(locked=False)
    entrada(wd.cell(l, 3), None)
    wd.row_dimensions[l].height = 22
ULTIMA = PRIMEIRA + len(itens) - 1
TOTAL = ULTIMA + 1
rotulo(wd.cell(TOTAL, 2), "CUSTO ESSENCIAL MENSAL")
wd.cell(TOTAL, 2).fill = fill(CREME)
saida(wd.cell(TOTAL, 3), f"=SUM(C{PRIMEIRA}:C{ULTIMA})", destaque=True)
wd.row_dimensions[TOTAL].height = 26
dicas = [
    "O que entra aqui?",
    "Só o que você NÃO pode deixar de pagar nem por um mês.",
    "Fica de fora: delivery, streaming, roupas, saídas,",
    "compras parceladas não essenciais.",
    "",
    "Os nomes das linhas \"Outro essencial\" podem ser",
    "editados: clique e escreva o que for.",
]
for i, t in enumerate(dicas):
    c = wd.cell(PRIMEIRA + i, 5, t)
    c.font = Font(size=9.5, bold=(i == 0), color=OURO_ESC if i == 0 else "5B645F")

# ---------------------------------------------------------------- 2 · Seu Número
wn = wb.create_sheet("2 · Seu Número")
base(wn, "Seu número", "Ajuste os campos amarelos. Todo o resto é calculado na hora.",
     {"A": 3, "B": 44, "C": 20, "D": 3, "E": 46})
faixa(wn, 5, "SEUS DADOS", "B:C")
rotulo(wn["B6"], "Custo essencial mensal (vem da aba 1)", negrito=False, cor="1B2622")
saida(wn["C6"], "='1 · Despesas'!C%d" % TOTAL)
rotulo(wn["B7"], "Quantos meses de reserva você quer?", negrito=False, cor="1B2622")
entrada(wn["C7"], 6, "0")
rotulo(wn["B8"], "Em quantos meses quer chegar lá? (prazo)", negrito=False, cor="1B2622")
entrada(wn["C8"], 18, "0")
rotulo(wn["B9"], "Quanto você já tem guardado hoje?", negrito=False, cor="1B2622")
entrada(wn["C9"], 0)
for l in range(6, 10):
    wn.row_dimensions[l].height = 24

faixa(wn, 11, "RESULTADO", "B:C")
rotulo(wn["B12"], "SEU NÚMERO (meta da reserva)")
saida(wn["C12"], "=C6*C7", destaque=True)
rotulo(wn["B13"], "Quanto ainda falta")
saida(wn["C13"], "=MAX(C12-C9,0)")
rotulo(wn["B14"], "GUARDAR POR MÊS")
saida(wn["C14"], "=IF(C8>0,C13/C8,0)", destaque=True)
rotulo(wn["B15"], "Guardar por semana (aprox.)")
saida(wn["C15"], "=C14*12/52")
rotulo(wn["B16"], "Guardar por dia (aprox.)")
saida(wn["C16"], "=C14*12/365")
rotulo(wn["B17"], "Você já tem quantos meses de reserva?")
saida(wn["C17"], "=IF(C6>0,C9/C6,0)", '0.0" meses"')
for l in range(12, 18):
    wn.row_dimensions[l].height = 26
wn.row_dimensions[12].height = 30
wn.row_dimensions[14].height = 30

faixa(wn, 19, "SEUS DEGRAUS", "B:C")
degraus = [
    ("Degrau 1 · R$ 1.000 ou meio mês (o menor)", "=MIN(1000,C6/2)"),
    ("Degrau 2 · 1 mês de custo essencial", "=C6*1"),
    ("Degrau 3 · 3 meses de custo essencial", "=C6*3"),
    ("Degrau 4 · reserva completa", "=C12"),
]
for i, (t, f) in enumerate(degraus):
    l = 20 + i
    rotulo(wn.cell(l, 2), t, negrito=False, cor="1B2622")
    saida(wn.cell(l, 3), f)
    wn.row_dimensions[l].height = 24

notas = [
    "Como ler",
    "Seu número = custo essencial × meses de reserva.",
    "Guardar por mês = (seu número − o que já tem) ÷ prazo.",
    "",
    "O valor mensal não cabe no seu orçamento?",
    "Aumente o prazo no campo amarelo e veja o novo valor.",
    "Ou compare vários prazos na aba 4 · Simulador de prazo.",
    "",
    "Os valores não consideram rendimento: na prática,",
    "o dinheiro aplicado rende e você chega um pouco antes.",
]
for i, t in enumerate(notas):
    c = wn.cell(6 + i, 5, t)
    c.font = Font(size=9.5, bold=t in ("Como ler", "O valor mensal não cabe no seu orçamento?"),
                  color=OURO_ESC if t in ("Como ler", "O valor mensal não cabe no seu orçamento?") else "5B645F")

dv_meses = DataValidation(type="whole", operator="between", formula1=1, formula2=24,
                          error="Use um número inteiro entre 1 e 24.", showErrorMessage=True)
dv_prazo = DataValidation(type="whole", operator="between", formula1=1, formula2=60,
                          error="Use um número inteiro de meses entre 1 e 60.", showErrorMessage=True)
dv_valor = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                          error="Use um valor igual ou maior que zero.", showErrorMessage=True)
wn.add_data_validation(dv_meses)
wn.add_data_validation(dv_prazo)
wn.add_data_validation(dv_valor)
dv_meses.add("C7")
dv_prazo.add("C8")
dv_valor.add("C9")

# ---------------------------------------------------------------- 3 · Acompanhamento
wa = wb.create_sheet("3 · Acompanhamento")
base(wa, "Acompanhamento mês a mês", "Todo mês, anote na coluna amarela quanto você depositou. Se precisou sacar, anote com sinal de menos.",
     {"A": 3, "B": 8, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 22})
wa["B4"] = "Já tinha guardado:"
wa["B4"].font = Font(size=10, color="5B645F")
wa.merge_cells("B4:C4")
saida(wa["D4"], "='2 · Seu Número'!C9")
wa["E4"] = "Meta final:"
wa["E4"].font = Font(size=10, color="5B645F")
wa["E4"].alignment = Alignment(horizontal="right")
saida(wa["F4"], "='2 · Seu Número'!C12")
cab = ["Mês", "Depositei no mês", "Total guardado", "Meta até aqui", "Diferença", "Quanto falta", "Situação"]
for i, t in enumerate(cab):
    c = wa.cell(6, 2 + i, t)
    c.fill = fill(VERDE)
    c.font = Font(size=9.5, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
wa.row_dimensions[6].height = 30
N = 'Seu Número'
for m in range(1, 61):
    l = 6 + m
    ativo = f"{m}<='2 · {N}'!$C$8"
    c = wa.cell(l, 2, m)
    c.alignment = Alignment(horizontal="center")
    c.font = Font(size=10, bold=True, color=VERDE)
    c.border = borda
    entrada(wa.cell(l, 3), None)
    saida(wa.cell(l, 4), f'=IF({ativo},$D$4+SUM($C$7:C{l}),"")')
    saida(wa.cell(l, 5), f"=IF({ativo},$D$4+'2 · {N}'!$C$14*B{l},\"\")")
    saida(wa.cell(l, 6), f'=IF({ativo},D{l}-E{l},"")')
    saida(wa.cell(l, 7), f'=IF({ativo},MAX($F$4-D{l},0),"")')
    s = wa.cell(l, 8, f'=IF({ativo},IF(C{l}="","",IF(D{l}>=$F$4,"Reserva completa!",'
                      f'IF(F{l}>=0,"No ritmo","Abaixo da meta"))),"")')
    s.border = borda
    s.alignment = Alignment(horizontal="center")
    s.font = Font(size=10, bold=True)
faixa_ultima = 6 + 60
wa.conditional_formatting.add(f"H7:H{faixa_ultima}", CellIsRule(operator="equal", formula=['"No ritmo"'],
                              font=Font(color="1E4A3D", bold=True), fill=fill("E4ECE7")))
wa.conditional_formatting.add(f"H7:H{faixa_ultima}", CellIsRule(operator="equal", formula=['"Abaixo da meta"'],
                              font=Font(color="8A3A2E", bold=True), fill=fill("F6E3DC")))
wa.conditional_formatting.add(f"H7:H{faixa_ultima}", CellIsRule(operator="equal", formula=['"Reserva completa!"'],
                              font=Font(color="FFFFFF", bold=True), fill=fill(OURO_ESC)))
# esconde visualmente as linhas além do prazo
wa.conditional_formatting.add(f"B7:C{faixa_ultima}", FormulaRule(formula=[f"$B7>'2 · {N}'!$C$8"],
                              font=Font(color="C8C8C8"), fill=fill("FFFFFF")))
wa.freeze_panes = "B7"

# ---------------------------------------------------------------- 4 · Simulador de prazo
wsim = wb.create_sheet("4 · Simulador de prazo")
base(wsim, "Simulador de prazo", "Quanto guardar por mês para chegar ao seu número em cada prazo. Usa os dados da aba 2.",
     {"A": 3, "B": 16, "C": 20, "D": 20, "E": 20})
for i, t in enumerate(["Prazo (meses)", "Guardar por mês", "Por semana (aprox.)", "Por dia (aprox.)"]):
    c = wsim.cell(5, 2 + i, t)
    c.fill = fill(VERDE)
    c.font = Font(size=9.5, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center")
wsim.row_dimensions[5].height = 24
for i, p in enumerate([6, 9, 12, 15, 18, 24, 30, 36]):
    l = 6 + i
    c = wsim.cell(l, 2, p)
    c.alignment = Alignment(horizontal="center")
    c.font = Font(size=11, bold=True, color=VERDE)
    c.border = borda
    saida(wsim.cell(l, 3), f"='2 · {N}'!$C$13/B{l}")
    saida(wsim.cell(l, 4), f"=C{l}*12/52")
    saida(wsim.cell(l, 5), f"=C{l}*12/365")
    wsim.row_dimensions[l].height = 22
wsim.conditional_formatting.add("B6:E13", FormulaRule(formula=[f"$B6='2 · {N}'!$C$8"],
                                fill=fill(ENTRADA), font=Font(bold=True, color=OURO_ESC)))
c = wsim["B15"]
c.value = "A linha destacada é o prazo que você escolheu na aba 2."
c.font = Font(size=9.5, color="5B645F")

# ---------------------------------------------------------------- proteção e acabamento
for w in wb.worksheets:
    w.protection.sheet = True
    w.protection.formatCells = False
    w.protection.formatColumns = False
    w.protection.formatRows = False
    w.sheet_properties.tabColor = OURO if w.title == "Comece aqui" else VERDE3
    w.page_setup.orientation = "portrait"
    w.page_setup.fitToWidth = 1
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToHeight = 0
wb.active = 0
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print("ok", SAIDA)
